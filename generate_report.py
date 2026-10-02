# -*- coding: utf-8 -*-
"""
長輩月活動評分與能力指標分析自動化轉檔程式（完整修正版）
- 自動適應表頭與評分欄位位移
- 第 1 頁綜合雷達圖：中心點為 0、移除 style 衝突，正常展開五角形
- 第 2 頁各類別獨立雷達圖：頂部標題置外不壓字、各端點標註「指標名稱與分數」
"""

import os
import re
import pandas as pd
import openpyxl
from openpyxl.chart import RadarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def clean_act_name(s):
    """清洗活動名稱，去除括弧與輔助註記"""
    s = re.sub(r'[\(\（].*?[\)\）]', '', str(s))
    return s.strip()


def extract_elder_info_smart(df):
    """智慧擷取長輩姓名、CMS 等級與評估月份"""
    info = {'姓名': '', 'CMS': '', '時間區間': ''}
    max_r = min(15, len(df))
    max_c = min(18, len(df.columns))

    for r in range(max_r):
        for c in range(max_c):
            val = str(df.iloc[r, c]).strip() if pd.notna(df.iloc[r, c]) else ''
            if not val or val == 'nan':
                continue

            # 抓取姓名
            if not info['姓名'] and any(k in val for k in ['個案姓名', '長輩姓名', '姓名']):
                m = re.search(r'姓名[：:\s]*([^\s,，\|｜\n]+)', val)
                if m and m.group(1):
                    info['姓名'] = m.group(1).replace('：', '').replace(':', '').strip()
                elif c + 1 < max_c:
                    next_val = str(df.iloc[r, c + 1]).strip()
                    if next_val and next_val != 'nan' and not any(k in next_val for k in ['CMS', '等級', '性別', '年齡', '身分證']):
                        info['姓名'] = next_val

            # 抓取 CMS 等級
            if not info['CMS'] and ('CMS' in val.upper() or '長照等級' in val or '失能等級' in val):
                m = re.search(r'(?:CMS|長照等級|失能等級|等級)[：:\s]*([0-9０-９]+)', val, re.IGNORECASE)
                if m:
                    info['CMS'] = m.group(1)
                elif c + 1 < max_c:
                    next_val = str(df.iloc[r, c + 1]).strip()
                    m_next = re.search(r'([0-9０-９]+)', next_val)
                    if m_next:
                        info['CMS'] = m_next.group(1)

            # 抓取評估月份
            if not info['時間區間']:
                m_date = re.search(r'(\d{4}[-/年]\d{1,2}(?:[-/月]\d{1,2})?|\d{2,3}\.\d{1,2})', val)
                if m_date:
                    info['時間區間'] = m_date.group(1)
                elif any(k in val for k in ['評估月份', '評估期間', '時間區間', '月份']):
                    if c + 1 < max_c:
                        next_val = str(df.iloc[r, c + 1]).strip()
                        m_next = re.search(r'(\d{4}[-/年]\d{1,2}|\d{2,3}\.\d{1,2})', next_val)
                        if m_next:
                            info['時間區間'] = m_next.group(1)

    if not info['姓名']:
        info['姓名'] = '個案'
    return info


def build_category_mapping(category_excel_path):
    """讀取活動清單對照表並建立分類字典與標準庫"""
    df_cat_raw = pd.read_excel(category_excel_path, sheet_name=0)
    cat_cols = [c for c in df_cat_raw.columns if not str(c).startswith('Unnamed')]
    mapping_dict = {}
    raw_category_rows = []
    for col in cat_cols:
        cat_name = str(col).strip()
        items = df_cat_raw[col].dropna().astype(str).tolist()
        for item in items:
            raw_item = item.strip()
            cleaned = clean_act_name(item)
            raw_category_rows.append({
                '原始活動名稱': raw_item,
                '正規化名稱': cleaned,
                '活動類型': cat_name
            })
            mapping_dict[raw_item] = cat_name
            if cleaned:
                mapping_dict[cleaned] = cat_name
    df_category_lib = pd.DataFrame(raw_category_rows).drop_duplicates(subset=['原始活動名稱'])
    return mapping_dict, df_category_lib


def process_monthly_scores(score_excel_path, mapping_dict):
    """解析評分紀錄並計算五大核心活動類型的指標均分（支援動態列識別）"""
    xls_score = pd.ExcelFile(score_excel_path)
    records = []
    elder_info = {'姓名': '', 'CMS': '', '時間區間': ''}

    for sheet in xls_score.sheet_names:
        df_s = pd.read_excel(score_excel_path, sheet_name=sheet, header=None)

        # 智慧擷取個案資訊
        extracted = extract_elder_info_smart(df_s)
        for k, v in extracted.items():
            if v and (not elder_info.get(k) or elder_info.get(k) == '個案'):
                elder_info[k] = v

        # 動態定位關鍵列的列號（自動適應表頭位移）
        date_row_idx = None
        am_act_idx = None
        pm_act_idx = None
        indicators_am = {}
        indicators_pm = {}

        for r_idx in range(min(35, len(df_s))):
            first_val = str(df_s.iloc[r_idx, 0]).strip()
            if '日期' in first_val and date_row_idx is None:
                date_row_idx = r_idx
            elif '上午活動名稱' in first_val or '上午活動' in first_val:
                am_act_idx = r_idx
            elif '下午活動名稱' in first_val or '下午活動' in first_val:
                pm_act_idx = r_idx
            elif am_act_idx is not None and pm_act_idx is None:
                for ind in ['參與度', '自我功能', '情緒控制力', '與他人互動', '肢體表現']:
                    if ind in first_val and ind not in indicators_am.values():
                        indicators_am[r_idx] = ind
            elif pm_act_idx is not None:
                for ind in ['參與度', '自我功能', '情緒控制力', '與他人互動', '肢體表現']:
                    if ind in first_val and ind not in indicators_pm.values():
                        indicators_pm[r_idx] = ind

        if date_row_idx is None or am_act_idx is None or pm_act_idx is None:
            continue

        row_date = df_s.iloc[date_row_idx]
        row_am_act = df_s.iloc[am_act_idx]
        row_pm_act = df_s.iloc[pm_act_idx]

        for col_idx in range(1, len(df_s.columns)):
            date_val = row_date.iloc[col_idx]
            if pd.isna(date_val) or str(date_val).strip() in ['累計', 'NaN', 'nan', '']:
                continue
            try:
                m_day = re.search(r'(\d+)', str(date_val))
                if not m_day:
                    continue
                day_num = int(m_day.group(1))
            except Exception:
                continue

            # 上午場次
            am_act = row_am_act.iloc[col_idx]
            if pd.notna(am_act) and str(am_act).strip() not in ['', 'nan', 'NaN']:
                act_str = str(am_act).strip()
                rec = {
                    '週次工作表': sheet,
                    '日期': day_num,
                    '時段': '上午',
                    '活動名稱': act_str,
                    '活動名稱_清洗': clean_act_name(act_str)
                }
                for r_i, ind_name in indicators_am.items():
                    rec[ind_name] = pd.to_numeric(df_s.iloc[r_i, col_idx], errors='coerce')
                records.append(rec)

            # 下午場次
            pm_act = row_pm_act.iloc[col_idx]
            if pd.notna(pm_act) and str(pm_act).strip() not in ['', 'nan', 'NaN']:
                act_str = str(pm_act).strip()
                rec = {
                    '週次工作表': sheet,
                    '日期': day_num,
                    '時段': '下午',
                    '活動名稱': act_str,
                    '活動名稱_清洗': clean_act_name(act_str)
                }
                for r_i, ind_name in indicators_pm.items():
                    rec[ind_name] = pd.to_numeric(df_s.iloc[r_i, col_idx], errors='coerce')
                records.append(rec)

    if not elder_info['姓名']:
        elder_info['姓名'] = '個案'

    df_records = pd.DataFrame(records)
    if df_records.empty:
        raise ValueError("未在檔案中找到任何有效的活動與評分資料，請確認工作表格式！")

    def assign_category(row):
        act = row['活動名稱']
        act_c = row['活動名稱_清洗']
        if act in mapping_dict:
            return mapping_dict[act]
        if act_c in mapping_dict:
            return mapping_dict[act_c]
        for k, v in mapping_dict.items():
            if k in act or act in k or k in act_c or act_c in k:
                return v
        if any(w in act for w in ['手作', '相框', '傘', '貼畫', '剪紙', '畫']):
            return '手部(精細)'
        if any(w in act for w in ['瑜珈', '肌不可失', '全身']):
            return '上下肢全身'
        if any(w in act for w in ['球', '拍球', '拋接']):
            return '手部(大肢體)'
        if any(w in act for w in ['踢', '下肢', '踏步']):
            return '下肢'
        if any(w in act for w in ['認知', '數', '猜', '懷舊', '連連看', '拼圖', '賓果']):
            return '認知'
        return '其他'

    df_records['活動類型'] = df_records.apply(assign_category, axis=1)

    ind_cols = ['參與度', '自我功能', '情緒控制力', '與他人互動', '肢體表現']
    core_five = ['手部(大肢體)', '手部(精細)', '下肢', '上下肢全身', '認知']

    summary_df = df_records.groupby('活動類型')[ind_cols].mean().round(1).reset_index()
    count_s = df_records.groupby('活動類型')['活動名稱'].count().reset_index().rename(columns={'活動名稱': '參與場次'})
    summary_df = pd.merge(summary_df, count_s, on='活動類型')

    summary_df['sort_key'] = summary_df['活動類型'].apply(lambda x: core_five.index(x) if x in core_five else 99)
    summary_df = summary_df.sort_values('sort_key').drop(columns=['sort_key']).reset_index(drop=True)

    overall_vals = {'活動類型': '【全月總平均】', '參與場次': len(df_records)}
    for col in ind_cols:
        overall_vals[col] = round(df_records[col].mean(), 1)
    summary_df = pd.concat([summary_df, pd.DataFrame([overall_vals])], ignore_index=True)

    return elder_info, df_records, summary_df, core_five


def export_excel_report(output_file, elder_info, df_records, summary_df, df_category_lib, core_five):
    """匯出格式化 Excel 報表：生成綜合表、獨立類別分析圖及分數表、明細與標準庫"""
    wb = openpyxl.Workbook()
    navy_header = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    navy_accent = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    light_blue = PatternFill(start_color="E9EEF4", end_color="E9EEF4", fill_type="solid")
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    white_font = Font(name="微軟正黑體", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="微軟正黑體", size=15, bold=True, color="1F4E79")
    subtitle_font = Font(name="微軟正黑體", size=10, color="595959")
    regular_font = Font(name="微軟正黑體", size=10)
    bold_font = Font(name="微軟正黑體", size=10, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )
    double_bottom_border = Border(
        left=Side(style='thin', color='E0E0E0'), right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='1F4E79'), bottom=Side(style='double', color='1F4E79')
    )

    # =========================================================================
    # 工作表 1: 活動指標總表與綜合分析
    # =========================================================================
    ws_dash = wb.active
    ws_dash.title = "活動指標總表與綜合分析"
    ws_dash.views.sheetView[0].showGridLines = True
    ws_dash['B2'] = "長輩各活動類型五大指標分析總表"
    ws_dash['B2'].font = title_font

    cms_str = f"CMS {elder_info.get('CMS')} 級" if elder_info.get('CMS') else "CMS (未填)"
    ws_dash['B3'] = f"評估月份：{elder_info.get('時間區間', '')} ｜ 長輩姓名：{elder_info.get('姓名', '')} ｜ 長照等級：{cms_str} ｜ 總活動場次：{len(df_records)} 場"
    ws_dash['B3'].font = subtitle_font

    headers = ['活動類型', '參與度', '自我功能', '情緒控制力', '與他人互動', '肢體表現', '參與場次']
    for c_idx, h in enumerate(headers, start=2):
        cell = ws_dash.cell(row=5, column=c_idx, value=h)
        cell.fill = navy_header
        cell.font = white_font
        cell.alignment = Alignment(horizontal="left" if c_idx == 2 else "center", vertical="center")
    ws_dash.row_dimensions[5].height = 26

    core_row_count = 0
    for r_idx, row in summary_df.iterrows():
        cur_row = 6 + r_idx
        is_overall = (r_idx == len(summary_df) - 1)
        if row['活動類型'] in core_five:
            core_row_count += 1

        vals = [
            row['活動類型'], row['參與度'], row['自我功能'], row['情緒控制力'],
            row['與他人互動'], row['肢體表現'], int(row['參與場次'])
        ]
        for c_idx, val in enumerate(vals, start=2):
            cell = ws_dash.cell(row=cur_row, column=c_idx, value=val)
            cell.font = bold_font if is_overall else regular_font
            cell.alignment = Alignment(horizontal="left" if c_idx == 2 else "center", vertical="center")
            cell.border = double_bottom_border if is_overall else thin_border
            if is_overall:
                cell.fill = light_blue
            elif r_idx % 2 == 1:
                cell.fill = zebra_fill
            if 3 <= c_idx <= 7:
                cell.number_format = '0.0'
        ws_dash.row_dimensions[cur_row].height = 20

    # 1. 綜合雷達圖（正常展開不壓縮）
    chart_comb = RadarChart()
    chart_comb.type = "standard"
    chart_comb.title = "各活動類型五大能力指標綜合比較"
    chart_comb.width = 18
    chart_comb.height = 13

    data_max_row = 5 + core_row_count
    data_comb = Reference(ws_dash, min_col=2, min_row=6, max_col=7, max_row=data_max_row)
    chart_comb.add_data(data_comb, titles_from_data=True, from_rows=True)

    cat_labels_dash = Reference(ws_dash, min_col=3, min_row=5, max_col=7, max_row=5)
    chart_comb.set_categories(cat_labels_dash)
    
    # 中心點設為 0.0，使蜘蛛網正常居中向外展開，刻度上限為 5.0
    chart_comb.y_axis.scaling.min = 0.0
    chart_comb.y_axis.scaling.max = 5.0

    marker_types = ["circle", "square", "triangle", "diamond", "star"]
    for idx, s in enumerate(chart_comb.series):
        s.marker = Marker(symbol=marker_types[idx % len(marker_types)], size=7)
        if s.graphicalProperties and s.graphicalProperties.line:
            s.graphicalProperties.line.width = 22000

    ws_dash.add_chart(chart_comb, "B16")

    # =========================================================================
    # 工作表 2: 各類別獨立分析（獨立標題置外、頂點標註指標名稱與分數）
    # =========================================================================
    ws_single = wb.create_sheet(title="各類別獨立分析")
    ws_single.views.sheetView[0].showGridLines = True
    ws_single['B2'] = "長輩各活動類別能力指標分析"
    ws_single['B2'].font = title_font
    ws_single['B3'] = f"評估月份：{elder_info.get('時間區間', '')} ｜ 長輩姓名：{elder_info.get('姓名', '')} ｜ 每類獨立呈現，附各指標具體分數"
    ws_single['B3'].font = subtitle_font

    indicators = ['參與度', '自我功能', '情緒控制力', '與他人互動', '肢體表現']

    layout_config = [
        {'chart_pos': 'B5',  'table_col': 10, 'table_row': 6},   # 左上
        {'chart_pos': 'N5',  'table_col': 22, 'table_row': 6},   # 右上
        {'chart_pos': 'B22', 'table_col': 10, 'table_row': 23},  # 左中
        {'chart_pos': 'N22', 'table_col': 22, 'table_row': 23},  # 右中
        {'chart_pos': 'B39', 'table_col': 10, 'table_row': 40},  # 左下
    ]

    for idx, cat_name in enumerate(core_five):
        match_row = None
        for r_num in range(6, 6 + len(summary_df)):
            if ws_dash.cell(r_num, 2).value == cat_name:
                match_row = r_num
                break
        if not match_row:
            continue

        cfg = layout_config[idx]
        t_col = cfg['table_col']
        t_row = cfg['table_row']

        # 將類別標題寫在圖表上方儲存格中（移開圖表內部，徹底避免壓到頂部分數）
        chart_col_letter = cfg['chart_pos'][0]
        chart_row_num = int(cfg['chart_pos'][1:])
        title_cell = ws_single.cell(row=chart_row_num - 1, column=ord(chart_col_letter) - ord('A') + 1, value=f"【{cat_name}】能力指標")
        title_cell.font = Font(name="微軟正黑體", size=12, bold=True, color="1F4E79")

        # 右側評分表格
        c_h1 = ws_single.cell(row=t_row, column=t_col, value="評估指標")
        c_h2 = ws_single.cell(row=t_row, column=t_col + 1, value=cat_name)
        for ch in [c_h1, c_h2]:
            ch.fill = navy_accent
            ch.font = white_font
            ch.alignment = Alignment(horizontal="center", vertical="center")
        ws_single.row_dimensions[t_row].height = 22

        for ind_idx, ind_name in enumerate(indicators):
            row_curr = t_row + 1 + ind_idx
            score_val = ws_dash.cell(row=match_row, column=3 + ind_idx).value

            c_name = ws_single.cell(row=row_curr, column=t_col, value=ind_name)
            c_score = ws_single.cell(row=row_curr, column=t_col + 1, value=score_val)

            c_name.font = regular_font
            c_name.alignment = Alignment(horizontal="left", vertical="center")
            c_name.border = thin_border

            c_score.font = bold_font
            c_score.alignment = Alignment(horizontal="center", vertical="center")
            c_score.border = thin_border
            c_score.number_format = '0.0'

            if ind_idx % 2 == 1:
                c_name.fill = zebra_fill
                c_score.fill = zebra_fill
            ws_single.row_dimensions[row_curr].height = 20

        # 雷達分析圖
        c_single = RadarChart()
        c_single.type = "standard"
        c_single.title = None  # 移除圖內置中標題，留出頂部完整空間
        c_single.width = 15.5
        c_single.height = 11.5
        c_single.y_axis.scaling.min = 0.0
        c_single.y_axis.scaling.max = 5.0

        cats_ref = Reference(ws_single, min_col=t_col, min_row=t_row + 1, max_row=t_row + len(indicators))
        data_ref = Reference(ws_single, min_col=t_col + 1, min_row=t_row, max_row=t_row + len(indicators))
        
        c_single.add_data(data_ref, titles_from_data=True, from_rows=False)
        c_single.set_categories(cats_ref)

        # 頂點同時標明「指標名稱」與「分數」
        c_single.dataLabels = DataLabelList()
        c_single.dataLabels.showVal = True         # 顯示分數
        c_single.dataLabels.showCatName = True     # 顯示指標名稱
        c_single.dataLabels.separator = "\n"       # 上下分行排列
        c_single.dataLabels.showSerName = False
        c_single.dataLabels.showLegendKey = False
        c_single.legend = None

        for s in c_single.series:
            s.marker = Marker(symbol=marker_types[idx % len(marker_types)], size=7)
            if s.graphicalProperties and s.graphicalProperties.line:
                s.graphicalProperties.line.width = 24000

        ws_single.add_chart(c_single, cfg['chart_pos'])

    # =========================================================================
    # 工作表 3: 全月活動明細與分類對照
    # =========================================================================
    ws_detail = wb.create_sheet(title="全月活動明細與分類對照")
    ws_detail.views.sheetView[0].showGridLines = True
    detail_headers = ['序號', '週次表單', '日期', '時段', '活動名稱', '自動對應活動類型', '參與度', '自我功能', '情緒控制力', '與他人互動', '肢體表現']
    for c_idx, h in enumerate(detail_headers, start=1):
        cell = ws_detail.cell(row=1, column=c_idx, value=h)
        cell.fill = navy_accent
        cell.font = white_font
        cell.alignment = Alignment(horizontal="left" if c_idx in [5, 6] else "center", vertical="center")
    ws_detail.row_dimensions[1].height = 22

    for idx, r in df_records.iterrows():
        row_num = idx + 2
        row_vals = [
            idx + 1, r['週次工作表'], f"{r['日期']}日", r['時段'], r['活動名稱'],
            r['活動類型'], r['參與度'], r['自我功能'], r['情緒控制力'], r['與他人互動'], r['肢體表現']
        ]
        for c_idx, v in enumerate(row_vals, start=1):
            cell = ws_detail.cell(row=row_num, column=c_idx, value=v)
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left" if c_idx in [5, 6] else "center", vertical="center")
            if idx % 2 == 1:
                cell.fill = zebra_fill
            if 7 <= c_idx <= 11:
                cell.number_format = '0.0'

    # =========================================================================
    # 工作表 4: 活動分類標準庫
    # =========================================================================
    ws_lib = wb.create_sheet(title="活動分類標準庫")
    ws_lib.views.sheetView[0].showGridLines = True
    lib_headers = ['序號', '原始活動名稱', '正規化名稱', '所屬活動類型']
    for c_idx, h in enumerate(lib_headers, start=1):
        cell = ws_lib.cell(row=1, column=c_idx, value=h)
        cell.fill = navy_header
        cell.font = white_font
        cell.alignment = Alignment(horizontal="left" if c_idx in [2, 3, 4] else "center", vertical="center")

    for idx, r in df_category_lib.reset_index(drop=True).iterrows():
        row_num = idx + 2
        row_vals = [idx + 1, r['原始活動名稱'], r['正規化名稱'], r['活動類型']]
        for c_idx, v in enumerate(row_vals, start=1):
            cell = ws_lib.cell(row=row_num, column=c_idx, value=v)
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left" if c_idx in [2, 3, 4] else "center", vertical="center")
            if idx % 2 == 1:
                cell.fill = zebra_fill

    # 自動調整欄寬
    for ws in [ws_dash, ws_single, ws_detail, ws_lib]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                w = sum(2 if ord(char) > 127 else 1 for char in val_str)
                if w > max_len:
                    max_len = w
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    ws_dash.column_dimensions['A'].width = 4
    ws_dash.column_dimensions['B'].width = 16

    ws_single.column_dimensions['A'].width = 3
    for col_l in ['B', 'C', 'D', 'E', 'F', 'G', 'H']:
        ws_single.column_dimensions[col_l].width = 6
    ws_single.column_dimensions['I'].width = 3
    ws_single.column_dimensions['J'].width = 14
    ws_single.column_dimensions['K'].width = 12

    ws_single.column_dimensions['L'].width = 3
    ws_single.column_dimensions['M'].width = 3
    for col_l in ['N', 'O', 'P', 'Q', 'R', 'S', 'T']:
        ws_single.column_dimensions[col_l].width = 6
    ws_single.column_dimensions['U'].width = 3
    ws_single.column_dimensions['V'].width = 14
    ws_single.column_dimensions['W'].width = 12

    wb.save(output_file)
    print(f"   [完成] 已產出報表: {output_file}")


# =============================================================================
# 主執行入口
# =============================================================================
if __name__ == '__main__':
    print('=' * 60)
    print('長輩月活動評分與指標分析自動化批次轉檔系統')
    print('=' * 60)

    cat_files = [f for f in os.listdir('.') if f.endswith('.xlsx') and '活動清單' in f and not f.startswith('~$')]
    if not cat_files:
        print('【錯誤】找不到包含「活動清單」的 Excel 檔案！請確認資料夾內是否有活動清單。')
        exit()

    cat_file = cat_files[0]
    print(f'--> 成功鎖定分類清單庫: {cat_file}')
    mapping_dict, df_category_lib = build_category_mapping(cat_file)

    pending = [
        f for f in os.listdir('.')
        if f.endswith('.xlsx')
        and not f.startswith('~$')
        and f != cat_file
        and '長輩活動評分與分析系統' not in f
    ]

    if not pending:
        print('\n【提示】目前資料夾中沒有待處理的長輩評分檔案。')
    else:
        print(f'\n共發現 {len(pending)} 個長輩紀錄檔案，開始處理...\n')
        for f in pending:
            try:
                info, recs, smry, cf = process_monthly_scores(f, mapping_dict)
                elder_name = re.sub(r'[\\/:*?"<>|]', '', str(info.get('姓名', '個案'))).strip()
                month_str = re.sub(r'[\\/:*?"<>|]', '', str(info.get('時間區間', '當月'))).strip()

                output_file = f'長輩活動評分與分析系統_{elder_name}_{month_str}.xlsx'
                export_excel_report(output_file, info, recs, smry, df_category_lib, cf)
            except Exception as e:
                print(f'   [失敗] 處理 {f} 失敗，錯誤原因: {e}')

        print('\n所有長輩評分表已全數批次處理完成！')